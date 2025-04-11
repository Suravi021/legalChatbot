import streamlit as st
from llama_index.core import VectorStoreIndex, SimpleDirectoryReader, Settings
from llama_index.core.prompts.prompts import PromptTemplate
from llama_index.embeddings.langchain import LangchainEmbedding
from llama_index.llms.langchain import LangChainLLM
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_ollama import OllamaLLM

#jinja
from docx import Document
import re
import os
import json



# Set up the assistant
system_prompt = """
You are a legal assistant specialized in contract drafting and analysis.
Follow this process when helping users:
1. First ask for the parties involved in the agreement
2. Inquire about the general scenario or purpose of the contract
3. Based on the information, identify the type of contract (e.g., employment, lease, sale, NDA)
4. Ask follow-up questions specific to that contract type
5. Provide relevant legal considerations for the identified contract type

Always maintain a professional, precise tone and ensure legal compliance.
"""


ollama_llm = OllamaLLM(model="llama3", temperature=0.2, system=system_prompt)
llm = LangChainLLM(llm=ollama_llm)
embed_model = LangchainEmbedding(HuggingFaceEmbeddings(model_name="sentence-transformers/all-mpnet-base-v2"))

Settings.llm = llm
Settings.embed_model = embed_model

documents = SimpleDirectoryReader("documents").load_data()
index = VectorStoreIndex.from_documents(documents)
query_engine = index.as_query_engine()

conversation_template = PromptTemplate(
    """
    You are a legal assistant helping with contracts.
    
    Conversation history:
    {context}
    
    User's latest message: {question}
    
    Based on the conversation so far, respond helpfully about the contract. 
    If this is the start of the conversation, ask about the purpose of the agreement between the parties.
    """
)

def extract_json_from_text(text):
    match = re.search(r"\{.*\}", text, re.DOTALL)
    return match.group() if match else "{}"

def extract_fields_from_context(context_text: str):
    prompt = f"""
### Instruction:
Extract the following fields from the text below and return them as a JSON object.

Fields:
- parties
- effective_date
- agreement_duration
- payment_terms
- special_conditions

If a field is missing, use null.

### Text:
\"\"\"{context_text}\"\"\"
### Response:
"""
    response = llm.complete(prompt)
    text_output = response.text
    json_part = extract_json_from_text(text_output)
    return json.loads(json_part)

def replace_placeholders_in_docx(doc_path, replacements):
    doc = Document(doc_path)
    pattern = r"\{\{\s*(\w+)\s*\}\}"

    for paragraph in doc.paragraphs:
        if re.search(pattern, paragraph.text):
            inline = paragraph.runs
            for i in range(len(inline)):
                text = inline[i].text
                for key, value in replacements.items():
                    text = re.sub(r"\{\{\s*" + re.escape(key) + r"\s*\}\}", value, text)
                inline[i].text = text
    return doc

def get_context():
    context = ""
    for sender, msg in st.session_state.messages:
        context += sender
        context += " "
        context += msg
        context += " "
    return context
        

    

# STREAMLIT UI
st.set_page_config(page_title="Legal Contract Assistant", layout="wide")

st.title("📜 Legal Contract Assistant")

if "context" not in st.session_state:
    st.session_state.context = ""
if "messages" not in st.session_state:
    st.session_state.messages = []
if "docx" not in st.session_state:
    st.session_state.docx = False
if "json_file" not in st.session_state:
    st.session_state.json_file = False

st.write("Chat with the assistant to discuss your contract. When you're done, click **Extract Fields**.")

user_input = st.chat_input("You: ", key="user_input")

if user_input:
    st.session_state.context += f"\nUser: {user_input}. make your response as short as possible. Only enquire about the details do no suggest anything"
    
    use_rag = any(keyword in user_input.lower() for keyword in [
        "legal", "requirement", "clause", "provision", "standard", "example", "template"
    ])
    
    if use_rag and query_engine:
        rag_response = query_engine.query(user_input)
        rag_context = f"Based on relevant legal information: {str(rag_response)}\n\n"
        enhanced_prompt = conversation_template.format(
            context=st.session_state.context + "\n" + rag_context,
            question=user_input
        )
        response_text = llm.complete(enhanced_prompt).text
    else:
        formatted_prompt = conversation_template.format(
            context=st.session_state.context,
            question=user_input
        )
        response_text = llm.complete(formatted_prompt).text
    
    st.session_state.context += f"\nAI: {response_text}"
    st.session_state.messages.append(("You", user_input))
    st.session_state.messages.append(("Bot", response_text))

# Display the chat history
for sender, msg in st.session_state.messages:
    with st.chat_message(sender):
        st.markdown(msg)
docx = False

# Button to extract fields
if st.button("📤 Extract Contract Fields"):
    if st.button("Employment Agreement"):
        docx = "employment_agreement.docx"
        json_file = "employment_agreement.json"
        
    if st.button("Paid internship Agreement"):
        docx = "paid_internship_agreement.docx"
        json_file = "paid_internship_agreement.json"

    if st.button("Model Partnership Agreement"):
        docx = "model_partnership_agreement.docx"
        json_file = "model_partnership_agreement.json"
    
    if st.button("Non Disclosure Agreement"):
        docx = "non_disclosure_agreement.docx"
        json_file = "non_disclosure_agreement.json"
    
if docx:
    print("yaaaaaaay")
    JSON_FILE = os.path.join("json_templates", json_file)
    with open(JSON_FILE, 'r') as f:
        fields = json.load(f)
    replacements = {}
    context = get_context()
    for field in fields:
        prompt = f"Generate a realistic value for: {field} with the context {context}, just the answer not anything else"
        temp = llm.complete(prompt=prompt).text 
        if temp:
            replacements[field] = temp
        else:
            prompt = f"Ask relavent question to know about the {field} in this context"
            temp = llm.complete(prompt=prompt).text
            

        print(f"{field} => {replacements[field]}")

    # docx = 
    updated_doc = replace_placeholders_in_docx(os.path.join("docs_words", docx), replacements)
    updated_doc.save(os.path.join("outputs", docx))
    print(f"Saved updated DOCX as {os.path.join("outputs", docx)}")

extracted = extract_fields_from_context(st.session_state.context)
st.subheader("📑 Extracted Contract Details (JSON):")
st.json(extracted)
